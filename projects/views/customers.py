import math
import random
from django.utils import timezone
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.db.models import Q, F as models_F, Count
from django.views.decorators.http import require_POST
from datetime import date, timedelta
import json

from ..models import Project, ProjectLog, Customer, Comment, Message, Notification, TeamMessage, StaffProfile, LeaveRequest, InstallationReport, ReportPhoto, SatisfactionNote, CustomerProfile, ProjectDocument, Product, PickingList, PickingListItem, PickingTemplate, PickingTemplateItem, MaterialPrice, ProjectCost, ProjectCostLine, UprightAccessory, AccessoryOverride, Reminder, FittingCrew, Supplier, PurchaseOrder, PurchaseOrderLine, StockMovement, FittingNote, ProjectQuote, PriceListItem, QuotePhoto, QuoteAttachedPhoto, ProformaInvoice, DeliveryPhase, ProjectPresence, CustomerContact, CustomerNote, CustomerDeliveryAddress
from ..forms import RegisterForm, ProjectForm


# ── Auth ──────────────────────────────────────────────────────────────────────

from django.contrib.auth.views import (
    PasswordResetView, PasswordResetDoneView,
    PasswordResetConfirmView, PasswordResetCompleteView,
)

@login_required
def customer_autocomplete(request):
    q = request.GET.get('q','').strip()
    if len(q) < 1:
        return JsonResponse([], safe=False)
    results = CustomerProfile.objects.filter(name__icontains=q).values('id', 'name')[:8]
    return JsonResponse(list(results), safe=False)




@login_required
def customer_history(request, pk):
    customer = get_object_or_404(CustomerProfile, pk=pk)
    projects = Project.objects.filter(
        customer__iexact=customer.name
    ).order_by('-created_at').select_related('assigned_to')
    total_projects = projects.count()
    # Total sell price from completed costings
    total_value = 0
    for p in projects:
        if hasattr(p, 'cost'):
            c = p.cost
            lines_total = sum(l.line_total for l in c.lines.all())
            uprights = sum(int(float(l.quantity))*2 for l in c.lines.all() if l.line_type=='frame')
            for acc in c.accessories.all():
                lines_total += float(acc.unit_price) * uprights
            markup_amt = lines_total * float(c.markup) / 100
            sell = lines_total + markup_amt + float(c.labour) + float(c.delivery)
            total_value += sell
    return render(request, 'projects/customer_history.html', {
        'customer': customer,
        'projects': projects,
        'total_projects': total_projects,
        'total_value': round(total_value, 2),
    })




@login_required
def customer_address_import(request):
    """Import addresses from a Sage 'Customer Address List' export — a
    different shape from the regular customer import: one row per
    address *line*, grouped into blocks under each customer's A/C row.
    Matches existing customers by account_number; doesn't create new
    customers. Names containing '***' (Sage's own dead/superseded
    account markers) get marked inactive rather than imported as live."""
    if not request.user.is_superuser:
        return JsonResponse({'error': 'Only administrators can run this import.'}, status=403)
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)

    import openpyxl, re
    f = request.FILES.get('file')
    if not f:
        return JsonResponse({'error': 'No file uploaded.'}, status=400)
    try:
        wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
        ws = wb.active
    except Exception as e:
        return JsonResponse({'error': f'Could not read file: {e}'}, status=400)

    postcode_re = re.compile(r'^[A-Z]{1,2}\d[A-Z\d]?\s*\d[A-Z]{2}$', re.IGNORECASE)

    # Find the header row by content rather than a fixed row number, since
    # Sage's report adds a few info rows above it that could shift.
    header_row_idx = None
    header = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        cells = [str(c).strip() if c is not None else '' for c in row]
        if 'A/C' in cells:
            header_row_idx = i
            header = cells
            break
    if not header_row_idx:
        return JsonResponse({'error': 'Could not find the header row — expected a column titled "A/C".'}, status=400)

    def col_idx(name):
        return header.index(name) if name in header else None

    i_ac = col_idx('A/C')
    i_addr = col_idx('Name & Address')
    if i_ac is None or i_addr is None:
        return JsonResponse({'error': 'Missing expected column(s): A/C, Name & Address.'}, status=400)

    def parse_block_addr(lines):
        """Turn a block's address lines into (addr1, addr2, town, county, postcode)
        based on the line count pattern found in the real export:
        5 lines = Addr1/Addr2/Town/County/Postcode, 4 = Addr1/Addr2/Town/Postcode,
        3 = Addr1/Town/Postcode. Non-UK addresses still use the last line as
        the postcode slot on a best-effort basis."""
        remaining = list(lines)
        postcode = remaining.pop() if remaining else ''
        addr1 = addr2 = town = county = ''
        n = len(remaining)
        if n >= 4:
            addr1, addr2, town, county = remaining[0], remaining[1], remaining[2], remaining[3]
        elif n == 3:
            addr1, addr2, town = remaining[0], remaining[1], remaining[2]
        elif n == 2:
            addr1, town = remaining[0], remaining[1]
        elif n == 1:
            addr1 = remaining[0]
        return addr1, addr2, town, county, postcode

    rows = list(ws.iter_rows(min_row=header_row_idx + 1, values_only=True))
    blocks = []
    current = None
    for row in rows:
        ac = row[i_ac] if i_ac < len(row) else None
        addr = row[i_addr] if i_addr < len(row) else None
        if ac is not None and str(ac).strip():
            if current:
                blocks.append(current)
            current = {'ac': str(ac).strip(), 'name': str(addr).strip() if addr else '', 'lines': []}
        elif addr is not None and str(addr).strip() and current:
            current['lines'].append(str(addr).strip())
    if current:
        blocks.append(current)

    matched = 0
    marked_inactive = 0
    not_matched = []
    to_update = []
    account_numbers = [b['ac'] for b in blocks]
    existing_by_ac = {
        c.account_number: c
        for c in CustomerProfile.objects.filter(account_number__in=account_numbers)
    }
    for b in blocks:
        customer = existing_by_ac.get(b['ac'])
        if not customer:
            not_matched.append(f"{b['ac']} — {b['name']}")
            continue
        addr1, addr2, town, county, postcode = parse_block_addr(b['lines'])
        customer.address_line1 = addr1
        customer.address_line2 = addr2
        customer.town = town
        customer.county = county
        customer.postcode = postcode
        if '***' in b['name']:
            customer.is_active = False
            marked_inactive += 1
        to_update.append(customer)
        matched += 1

    CustomerProfile.objects.bulk_update(
        to_update, ['address_line1', 'address_line2', 'town', 'county', 'postcode', 'is_active'], batch_size=500
    )

    return JsonResponse({
        'ok': True,
        'total_blocks': len(blocks),
        'matched': matched,
        'marked_inactive': marked_inactive,
        'not_matched_count': len(not_matched),
        'not_matched_sample': not_matched[:40],
    })


@login_required
def customer_import(request):
    if not request.user.is_superuser:
        return JsonResponse({'error': 'Only Tasos can import customers from Excel.'}, status=403)
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)
    import openpyxl
    f = request.FILES.get('file')
    if not f:
        return JsonResponse({'error': 'No file'}, status=400)
    try:
        wb = openpyxl.load_workbook(f)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return JsonResponse({'error': 'File is empty.'}, status=400)

        # Map columns by header name — robust against the spreadsheet's
        # column order changing, unlike matching by fixed position.
        header_row = [str(h).strip() if h is not None else '' for h in rows[0]]
        required = ['A/C', 'Name', 'Contact', 'Telephone',
                    'VAT Number', 'Company Reg. Number', 'Email 1', 'Email 3']
        missing = [h for h in required if h not in header_row]
        if missing:
            return JsonResponse({'error': f'Missing expected column(s): {", ".join(missing)}'}, status=400)
        col = {h: header_row.index(h) for h in required}

        def text(row, key):
            v = row[col[key]]
            return str(v).strip() if v is not None else ''

        created = 0
        skipped = 0
        for row in rows[1:]:
            name = text(row, 'Name')
            if not name:
                skipped += 1
                continue

            # Always create a new record — never match against, or update,
            # an existing one. Any duplicates this creates are intentional;
            # they get reconciled by hand afterwards, not merged here.
            CustomerProfile.objects.create(
                name=name,
                account_number=text(row, 'A/C'),
                contact_name=text(row, 'Contact'),
                phone=text(row, 'Telephone'),
                vat_number=text(row, 'VAT Number'),
                company_reg_number=text(row, 'Company Reg. Number'),
                email=text(row, 'Email 1'),
                email3=text(row, 'Email 3'),
            )
            created += 1
        return JsonResponse({'ok': True, 'created': created, 'skipped': skipped})
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=400)


@login_required
def customer_list(request):
    q = request.GET.get('q', '').strip()
    show_inactive = request.GET.get('show_inactive', '') == '1'
    letter = request.GET.get('letter', '').strip().upper()[:1]

    base = CustomerProfile.objects.all()
    if not show_inactive:
        base = base.filter(is_active=True)

    if q:
        # A text search overrides the index filter — show matches regardless
        # of which letter/number they start with.
        letter = ''
        customers = base.filter(
            Q(name__icontains=q) | Q(contact_name__icontains=q) |
            Q(email__icontains=q) | Q(phone__icontains=q)
        ).order_by('name')
    else:
        if not letter:
            letter = '0'
        customers = base.filter(name__istartswith=letter).order_by('name')

    # Per-letter/number counts for the index column — one query, not 36.
    from django.db.models.functions import Upper, Left
    from django.db.models import Count
    counts_qs = base.annotate(first_char=Upper(Left('name', 1))).values('first_char').annotate(n=Count('id'))
    counts = {row['first_char']: row['n'] for row in counts_qs}
    index_chars = [{'char': c, 'count': counts.get(c, 0)} for c in list('0123456789') + list('ABCDEFGHIJKLMNOPQRSTUVWXYZ')]

    # Paginate within the selected letter/search — groups are naturally much
    # smaller than the full list, so a modest page size keeps things fast
    # without needing many pages.
    from django.core.paginator import Paginator
    total_matching = customers.count()
    paginator = Paginator(customers, 50)
    page_num = request.GET.get('page', 1)
    try:
        page_obj = paginator.page(page_num)
    except Exception:
        page_obj = paginator.page(1)
    customers = page_obj.object_list

    # Page numbers to show at the bottom — all of them if there aren't too
    # many, otherwise a truncated range around the current page (with a
    # None entry meaning "show an ellipsis here") so jumping from page 1 to
    # page 11 of 12 doesn't take ten clicks through Next.
    num_pages = paginator.num_pages
    current = page_obj.number
    if num_pages <= 15:
        page_range = list(range(1, num_pages + 1))
    else:
        pages = {1, num_pages, current}
        for d in (1, 2):
            pages.add(current - d)
            pages.add(current + d)
        pages = sorted(p for p in pages if 1 <= p <= num_pages)
        page_range = []
        prev = None
        for p in pages:
            if prev is not None and p - prev > 1:
                page_range.append(None)
            page_range.append(p)
            prev = p

    return render(request, 'projects/customer_list.html', {
        'customers': customers, 'query': q,
        'show_inactive': show_inactive,
        'inactive_count': CustomerProfile.objects.filter(is_active=False).count(),
        'page_obj': page_obj, 'total_matching': total_matching,
        'letter': letter, 'index_chars': index_chars, 'page_range': page_range,
    })




@login_required
def customer_detail(request, pk):
    customer = get_object_or_404(CustomerProfile, pk=pk)
    if request.method == 'POST':
        old_name = customer.name
        customer.name         = request.POST.get('name', '').strip()
        customer.contact_name = request.POST.get('contact_name', '').strip()
        customer.email        = request.POST.get('email', '').strip()
        customer.email2       = request.POST.get('email2', '').strip()
        customer.email3       = request.POST.get('email3', '').strip()
        customer.phone        = request.POST.get('phone', '').strip()
        customer.vat_number   = request.POST.get('vat_number', '').strip()
        customer.company_reg_number = request.POST.get('company_reg_number', '').strip()
        customer.is_active = request.POST.get('is_active') == 'on'
        customer.eori_number  = request.POST.get('eori_number', '').strip()
        customer.account_number = request.POST.get('account_number', '').strip()
        customer.address_line1 = request.POST.get('address_line1', '').strip()
        customer.address_line2 = request.POST.get('address_line2', '').strip()
        customer.town          = request.POST.get('town', '').strip()
        customer.county        = request.POST.get('county', '').strip()
        customer.postcode      = request.POST.get('postcode', '').strip()
        customer.country       = request.POST.get('country', '').strip() or 'United Kingdom'
        customer.notes           = request.POST.get('notes', '').strip()
        customer.important_notes = request.POST.get('important_notes', '').strip()
        customer.save()
        # Cascade name change to all projects that reference the old name
        if customer.name and customer.name != old_name:
            affected = Project.objects.filter(customer__iexact=old_name)
            updated = 0
            for p in affected:
                p.customer = customer.name
                p.project_name = f"{customer.name} — {p.location}" if p.location else customer.name
                p.save(update_fields=['customer', 'project_name'])
                updated += 1
            if updated:
                messages.success(request, f'Customer updated. {updated} project{"s" if updated != 1 else ""} also updated to "{customer.name}".')
            else:
                messages.success(request, 'Customer updated.')
        else:
            messages.success(request, 'Customer updated.')
        return redirect('customer_detail', pk=pk)
    linked_projects = customer.projects.all()
    text_matched = Project.objects.filter(customer__iexact=customer.name).exclude(pk__in=linked_projects.values('pk'))
    projects = (linked_projects | text_matched).order_by('-created_at')
    pos = PurchaseOrder.objects.filter(project__in=projects).select_related('supplier', 'project').order_by('-created_at')
    contacts = customer.contacts.all()
    delivery_addresses = customer.delivery_addresses.select_related('source_project').all()
    relationship_notes = customer.relationship_notes.select_related('created_by').all()

    # ── Timeline: projects, quotes, notes, and installs, newest first ──────
    from django.urls import reverse
    timeline = []
    for p in projects:
        timeline.append({'date': p.created_at, 'kind': 'project', 'icon': '📁',
            'title': f'Project created: {p.project_name}', 'url': reverse('project_edit', args=[p.pk]), 'project': p})
        if p.status == 'cancelled' and p.lost_reason:
            timeline.append({'date': p.updated_at, 'kind': 'lost', 'icon': '✕',
                'title': f'Lost: {p.project_name}', 'detail': p.lost_reason, 'url': reverse('project_edit', args=[p.pk]), 'project': p})
    quotes = ProjectQuote.objects.filter(project__in=projects).select_related('project')
    for q in quotes:
        timeline.append({'date': q.updated_at, 'kind': 'quote', 'icon': '📄',
            'title': f'Quote for {q.project.project_name}', 'detail': f'£{q.main_price:,.2f}' if q.main_price else None, 'url': reverse('project_edit', args=[q.project.pk]), 'project': q.project})
    installs = InstallationReport.objects.filter(project__in=projects).select_related('project')
    for i in installs:
        timeline.append({'date': i.created_at, 'kind': 'install', 'icon': '🔧',
            'title': f'Installed: {i.project.project_name}', 'url': reverse('project_edit', args=[i.project.pk]), 'project': i.project})
    for n in relationship_notes:
        timeline.append({'date': n.created_at, 'kind': 'note', 'icon': '📝',
            'title': n.created_by.get_full_name() if n.created_by and n.created_by.get_full_name() else (n.created_by.username if n.created_by else 'Note'),
            'detail': n.text})
    timeline.sort(key=lambda e: e['date'], reverse=True)

    from ..countries import COUNTRIES
    return render(request, 'projects/customer_detail.html', {
        'customer': customer, 'projects': projects, 'pos': pos, 'countries': COUNTRIES,
        'contacts': contacts, 'relationship_notes': relationship_notes, 'timeline': timeline,
        'delivery_addresses': delivery_addresses,
    })


@login_required
@require_POST
def customer_contact_add(request, pk):
    customer = get_object_or_404(CustomerProfile, pk=pk)
    data = json.loads(request.body)
    name = (data.get('name') or '').strip()
    if not name:
        return JsonResponse({'error': 'Name is required.'}, status=400)
    next_order = (customer.contacts.count())
    contact = CustomerContact.objects.create(
        customer=customer, name=name,
        role=(data.get('role') or '').strip(),
        phone=(data.get('phone') or '').strip(),
        email=(data.get('email') or '').strip(),
        sort_order=next_order,
    )
    return JsonResponse({'ok': True, 'id': contact.pk})


@login_required
@require_POST
def customer_contact_update(request, pk):
    contact = get_object_or_404(CustomerContact, pk=pk)
    data = json.loads(request.body)
    if 'name' in data:
        name = (data.get('name') or '').strip()
        if not name:
            return JsonResponse({'error': 'Name is required.'}, status=400)
        contact.name = name
    if 'role' in data:
        contact.role = (data.get('role') or '').strip()
    if 'phone' in data:
        contact.phone = (data.get('phone') or '').strip()
    if 'email' in data:
        contact.email = (data.get('email') or '').strip()
    contact.save()
    return JsonResponse({'ok': True})


@login_required
@require_POST
def customer_contact_delete(request, pk):
    CustomerContact.objects.filter(pk=pk).delete()
    return JsonResponse({'ok': True})


@login_required
@require_POST
def customer_note_add(request, pk):
    customer = get_object_or_404(CustomerProfile, pk=pk)
    data = json.loads(request.body)
    text = (data.get('text') or '').strip()
    if not text:
        return JsonResponse({'error': 'Note text is required.'}, status=400)
    note = CustomerNote.objects.create(customer=customer, text=text, created_by=request.user)
    return JsonResponse({
        'ok': True, 'id': note.pk,
        'author': request.user.get_full_name() or request.user.username,
        'created_at': note.created_at.strftime('%d %b %Y, %H:%M'),
    })


@login_required
@require_POST
def customer_note_delete(request, pk):
    note = get_object_or_404(CustomerNote, pk=pk)
    if note.created_by_id and note.created_by_id != request.user.id and not request.user.is_superuser:
        return JsonResponse({'error': "You can only delete your own notes."}, status=403)
    note.delete()
    return JsonResponse({'ok': True})


@login_required
@require_POST
def customer_delivery_address_add(request, pk):
    customer = get_object_or_404(CustomerProfile, pk=pk)
    data = json.loads(request.body)
    next_order = customer.delivery_addresses.count()
    addr = CustomerDeliveryAddress.objects.create(
        customer=customer,
        line1=(data.get('line1') or '').strip(),
        line2=(data.get('line2') or '').strip(),
        city=(data.get('city') or '').strip(),
        county=(data.get('county') or '').strip(),
        postcode=(data.get('postcode') or '').strip(),
        country=(data.get('country') or '').strip() or 'United Kingdom',
        fao=(data.get('fao') or '').strip(),
        phone=(data.get('phone') or '').strip(),
        sort_order=next_order,
    )
    return JsonResponse({'ok': True, 'id': addr.pk})


@login_required
@require_POST
def customer_delivery_address_update(request, pk):
    addr = get_object_or_404(CustomerDeliveryAddress, pk=pk)
    data = json.loads(request.body)
    for field in ['line1', 'line2', 'city', 'county', 'postcode', 'country', 'fao', 'phone']:
        if field in data:
            setattr(addr, field, (data.get(field) or '').strip())
    addr.save()
    return JsonResponse({'ok': True})


@login_required
@require_POST
def customer_delivery_address_delete(request, pk):
    CustomerDeliveryAddress.objects.filter(pk=pk).delete()
    return JsonResponse({'ok': True})




@login_required
@require_POST
def customer_create(request):
    data = json.loads(request.body)
    name = data.get('name', '').strip()
    if not name:
        return JsonResponse({'error': 'Name required'}, status=400)
    try:
        existing = CustomerProfile.objects.filter(name__iexact=name).first()
        if existing:
            return JsonResponse({'id': existing.pk, 'name': existing.name, 'created': False})
        c = CustomerProfile.objects.create(
            name=name,
            contact_name=data.get('contact_name','').strip(),
            email=data.get('email','').strip(),
            phone=data.get('phone','').strip(),
            address=data.get('address','').strip(),
            notes=data.get('notes','').strip(),
        )
        return JsonResponse({'id': c.pk, 'name': c.name, 'created': True})
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=400)




@login_required
@require_POST
def customer_delete(request, pk):
    customer = get_object_or_404(CustomerProfile, pk=pk)
    customer.delete()
    return JsonResponse({'ok': True})


# ── Suppliers ─────────────────────────────────────────────────────────────────

